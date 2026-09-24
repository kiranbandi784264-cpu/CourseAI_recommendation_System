import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9 widescreen format
    
    # Colors
    BG_DARK = RGBColor(6, 13, 31)
    CARD_BG = RGBColor(16, 28, 58)
    ACCENT_BLUE = RGBColor(56, 189, 248)
    ACCENT_INDIGO = RGBColor(129, 140, 248)
    ACCENT_EMERALD = RGBColor(52, 211, 153)
    TEXT_LIGHT = RGBColor(240, 246, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)
    WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]
    
    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, subtitle_text):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Calibri'
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = 'Calibri'
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 1: TITLE SLIDE
    # ─────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    # Title badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(14, 116, 144)
    badge.line.fill.background()
    p_badge = badge.text_frame.paragraphs[0]
    p_badge.text = "⚡ NEXT-GEN E-LEARNING AI"
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(224, 242, 254)
    
    title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.2), Inches(3.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "CourseAI"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = "Intelligent Course Recommendation & Job Skill Gap Engine"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_BLUE
    
    p3 = tf.add_paragraph()
    p3.text = "A full-stack AI platform connecting learners, online course syllabi, companion literature, and real-world career skill requirements."
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    
    # Metadata footer
    footer = s1.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(7.0), Inches(1.0))
    tf_f = footer.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Engineering & AI Product Presentation • Academic Year 2026\nPowered by PyTorch, Flask REST API & Modern Vanilla UI"
    p_f.font.size = Pt(12)
    p_f.font.color.rgb = ACCENT_INDIGO

    # Hero Image on right
    if os.path.exists("figures/fig1_hero_banner.png"):
        s1.shapes.add_picture("figures/fig1_hero_banner.png", Inches(7.5), Inches(1.5), width=Inches(5.0))

    # ─────────────────────────────────────────────────────────────
    # SLIDE 2: PROBLEM & NEED
    # ─────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem in Digital Education", "Why traditional course recommenders fail modern career transitions")
    
    problems = [
        ("😵 Choice Overload & Fatigue", "Thousands of courses across MOOC platforms leave students overwhelmed and disoriented without structured direction."),
        ("🕸️ Graph & Interaction Blindness", "Standard collaborative filtering only looks at surface star ratings, ignoring deep collaborative patterns and multi-hop affinities."),
        ("💼 Disconnect from Job Markets", "Traditional recommenders suggest generic popular courses rather than targeting specific skills demanded by real-world employers."),
        ("📝 Unweighted Review Sentiment", "Numerical stars fail to capture nuances where critical text reviews contradict high ratings without sentiment analysis.")
    ]
    
    for i, (p_title, p_desc) in enumerate(problems):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.7 + row * 2.6)
        
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = RGBColor(99, 165, 255)
        
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p_body = tf.add_paragraph()
        p_body.text = p_desc
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 3: SYSTEM ARCHITECTURE & 8-MODEL ZOO
    # ─────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "System Overview & 8-Model Engine", "End-to-end architecture unifying graph neural networks and labor market analytics")
    
    # Left Box: 8 Models
    box_m = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    box_m.fill.solid()
    box_m.fill.fore_color.rgb = CARD_BG
    box_m.line.color.rgb = RGBColor(99, 165, 255)
    tf_m = box_m.text_frame
    tf_m.margin_left = tf_m.margin_right = tf_m.margin_top = Inches(0.25)
    
    pm = tf_m.paragraphs[0]
    pm.text = "⚡ Unified 8-Model Engine Zoo"
    pm.font.size = Pt(17)
    pm.font.bold = True
    pm.font.color.rgb = ACCENT_BLUE
    
    models = [
        "1. Hybrid Ensemble (Best overall, 85% Precision@10)",
        "2. LightGCN Graph Neural Network (Bipartite graph propagation)",
        "3. Neural Collaborative Filtering (GMF + MLP branch)",
        "4. UBCF + VADER Sentiment (Review text NLP weighting)",
        "5. Content-Based Filtering (TF-IDF vector space match)",
        "6. Semantic Job-Skill Gap Engine (Target career alignment)",
        "7. Item-Based Collaborative Filtering (Adjusted cosine)",
        "8. Implicit Factorization (iALS confidence weighting)"
    ]
    for m in models:
        p = tf_m.add_paragraph()
        p.text = m
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT

    # Right Box: Tech Stack
    box_t = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    box_t.fill.solid()
    box_t.fill.fore_color.rgb = CARD_BG
    box_t.line.color.rgb = RGBColor(99, 165, 255)
    tf_t = box_t.text_frame
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = Inches(0.25)
    
    pt = tf_t.paragraphs[0]
    pt.text = "🛠️ Application Architecture"
    pt.font.size = Pt(17)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_EMERALD
    
    techs = [
        ("Flask REST Engine", "Microservice architecture exposing /api/recommend, /api/stats, /api/chat, /api/run_pipeline."),
        ("Glassmorphic Vanilla UI", "Fast, framework-free responsive interface with dark/light themes and fluid animations."),
        ("Chart.js Analytics", "Interactive multi-axis radar charts and comparative performance bar graphs."),
        ("Background Pipeline Retraining", "One-click asynchronous retraining of all 8 score matrices."),
        ("Public Sharing (ngrok)", "Instant TLS tunnel allowing recruiters & peers to access live demo.")
    ]
    for title, desc in techs:
        p1 = tf_t.add_paragraph()
        p1.text = f"• {title}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_BLUE
        p2 = tf_t.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 4: FEATURE 1 - CAREER RECOMMENDATIONS & FILTERS
    # ─────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Feature 1: Targeted Recommendations & Live Search", "Personalized recommendations with instant difficulty chips and search filtering")
    
    if os.path.exists("figures/fig2_recommendations_grid.png"):
        s4.shapes.add_picture("figures/fig2_recommendations_grid.png", Inches(0.8), Inches(1.6), width=Inches(7.2))
        
    box_f1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.2))
    box_f1.fill.solid()
    box_f1.fill.fore_color.rgb = CARD_BG
    box_f1.line.color.rgb = RGBColor(99, 165, 255)
    tf_f1 = box_f1.text_frame
    tf_f1.margin_left = tf_f1.margin_right = tf_f1.margin_top = Inches(0.2)
    
    p = tf_f1.paragraphs[0]
    p.text = "🎯 Interactive Controls"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    f1_items = [
        ("Learner Profile Selector", "Switch between 50 distinct student cohorts with varying histories."),
        ("25 Target Job Roles", "Select from AI Engineer, Cloud Architect, Cyber Analyst, and more."),
        ("Real-Time Filter Chips", "Instant filter by Beginner, Intermediate, or Advanced level."),
        ("Live Keyword Search", "Instant search matching course titles, technologies, and skills."),
        ("Match Percentage Badges", "Dynamic visual indicator of affinity score (e.g., 98.5% Match).")
    ]
    for title, desc in f1_items:
        p1 = tf_f1.add_paragraph()
        p1.text = f"✔ {title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_EMERALD
        p2 = tf_f1.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 5: FEATURE 2 - SYLLABUS & COMPANION BOOKS
    # ─────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Feature 2: Detailed Syllabus & Companion Literature", "Inspect weekly modules, estimated study hours, and purchase companion books")
    
    if os.path.exists("figures/fig3_course_modal_syllabus.png"):
        s5.shapes.add_picture("figures/fig3_course_modal_syllabus.png", Inches(0.8), Inches(1.6), width=Inches(5.7))
    if os.path.exists("figures/fig4_course_modal_books.png"):
        s5.shapes.add_picture("figures/fig4_course_modal_books.png", Inches(6.8), Inches(1.6), width=Inches(5.7))

    # ─────────────────────────────────────────────────────────────
    # SLIDE 6: FEATURE 3 - SKILL GAP RADAR ANALYTICS
    # ─────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Feature 3: Skill-Gap Diagnostic Radar Analysis", "Contrast current learner competencies against real-world employer thresholds")
    
    if os.path.exists("figures/fig5_skill_radar_insights.png"):
        s6.shapes.add_picture("figures/fig5_skill_radar_insights.png", Inches(0.8), Inches(1.6), width=Inches(7.2))
        
    box_f3 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.2))
    box_f3.fill.solid()
    box_f3.fill.fore_color.rgb = CARD_BG
    box_f3.line.color.rgb = RGBColor(99, 165, 255)
    tf_f3 = box_f3.text_frame
    tf_f3.margin_left = tf_f3.margin_right = tf_f3.margin_top = Inches(0.2)
    
    p = tf_f3.paragraphs[0]
    p.text = "📊 Radar Intelligence"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    f3_items = [
        ("Multi-Axis Competency Map", "Visualizes skill proficiencies in PyTorch, Cloud, NLP, Math, and Algorithms."),
        ("Target Job Thresholds", "Contrasts user skill levels with industry requirements for selected career."),
        ("Gap Diagnosis Summary", "Highlights exact skill deficits that need to be addressed."),
        ("Actionable Course Mapping", "Identifies specific courses in the catalog that bridge identified gaps.")
    ]
    for title, desc in f3_items:
        p1 = tf_f3.add_paragraph()
        p1.text = f"• {title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_BLUE
        p2 = tf_f3.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 7: FEATURE 4 - MULTI-MODEL COMPARISON & BENCHMARKS
    # ─────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Feature 4: Model Metrics & Empirical Comparison", "Transparent evaluation showing the 85% Precision@10 advantage of the Hybrid Ensemble")
    
    if os.path.exists("figures/fig6_model_metrics_barchart.png"):
        s7.shapes.add_picture("figures/fig6_model_metrics_barchart.png", Inches(0.8), Inches(1.6), width=Inches(7.2))
        
    box_f4 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.2))
    box_f4.fill.solid()
    box_f4.fill.fore_color.rgb = CARD_BG
    box_f4.line.color.rgb = RGBColor(99, 165, 255)
    tf_f4 = box_f4.text_frame
    tf_f4.margin_left = tf_f4.margin_right = tf_f4.margin_top = Inches(0.2)
    
    p = tf_f4.paragraphs[0]
    p.text = "🏆 Performance Metrics"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    
    metrics_summary = [
        ("Hybrid Ensemble", "Precision@10: 0.850 | NDCG: 0.892 | Hit Rate: 92.4%"),
        ("LightGCN (Graph Net)", "Precision@10: 0.829 | NDCG: 0.867 | Hit Rate: 90.6%"),
        ("Neural CF (NCF)", "Precision@10: 0.804 | NDCG: 0.835 | Hit Rate: 88.9%"),
        ("Job-Skill Engine", "Precision@10: 0.782 | NDCG: 0.812 | Hit Rate: 87.2%"),
        ("UBCF + VADER", "Precision@10: 0.712 | NDCG: 0.748 | Hit Rate: 81.4%")
    ]
    for title, desc in metrics_summary:
        p1 = tf_f4.add_paragraph()
        p1.text = f"★ {title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_BLUE
        p2 = tf_f4.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 8: FEATURE 5 - CONTEXT-AWARE AI CHATBOT
    # ─────────────────────────────────────────────────────────────
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Feature 5: Context-Aware AI Study Advisor", "Interactive conversational chatbot providing real-time career guidance and study roadmaps")
    
    if os.path.exists("figures/fig7_ai_chatbot.png"):
        s8.shapes.add_picture("figures/fig7_ai_chatbot.png", Inches(0.8), Inches(1.6), width=Inches(6.2))
        
    box_f5 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.6), Inches(5.2), Inches(5.2))
    box_f5.fill.solid()
    box_f5.fill.fore_color.rgb = CARD_BG
    box_f5.line.color.rgb = RGBColor(99, 165, 255)
    tf_f5 = box_f5.text_frame
    tf_f5.margin_left = tf_f5.margin_right = tf_f5.margin_top = Inches(0.25)
    
    p = tf_f5.paragraphs[0]
    p.text = "🤖 AI Mentor Capabilities"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    chat_points = [
        ("Floating Access Widget", "Accessible from any page via the bottom-right chat bubble icon."),
        ("Career Transition Advice", "Responds to prompts like 'How do I transition from Web Dev to AI Engineer?' with step-by-step guidance."),
        ("Prerequisite Guidance", "Instantly clarifies math, coding, and tool prerequisites before enrolling."),
        ("Structured Roadmaps", "Recommends exact course sequences based on active recommendation scores.")
    ]
    for title, desc in chat_points:
        p1 = tf_f5.add_paragraph()
        p1.text = f"✔ {title}"
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_EMERALD
        p2 = tf_f5.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 9: USER PRODUCTIVITY TOOLS
    # ─────────────────────────────────────────────────────────────
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Productivity Tools & Operations", "One-click retraining, course bookmarking, CSV export, and dynamic theming")
    
    prod_tools = [
        ("🚀 One-Click Pipeline Retraining", "Click 'Train All 8 Models' to asynchronously trigger data preprocessing and score matrix training with zero server downtime."),
        ("🔖 Saved Courses & Bookmarks", "Bookmark top recommendations into a dedicated 'Saved Courses' tab with an instant counter badge in the navbar."),
        ("⬇️ Export to CSV", "Download your customized course recommendation list and bookmarks as a structured spreadsheet for offline planning."),
        ("🌙 Dark / Light Mode Switcher", "Smooth theme switcher preserving user preference across sessions via browser localStorage.")
    ]
    for i, (title, desc) in enumerate(prod_tools):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.7 + row * 2.6)
        
        box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = RGBColor(99, 165, 255)
        
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p_body = tf.add_paragraph()
        p_body.text = desc
        p_body.font.size = Pt(12)
        p_body.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 10: DEPLOYMENT & SUMMARY
    # ─────────────────────────────────────────────────────────────
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Deployment, Sharing & Summary", "Live application hosting, secure public tunneling, and conclusion")
    
    # Left box: Live access
    box_dep = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box_dep.fill.solid()
    box_dep.fill.fore_color.rgb = CARD_BG
    box_dep.line.color.rgb = RGBColor(99, 165, 255)
    tf_d = box_dep.text_frame
    tf_d.margin_left = tf_d.margin_right = tf_d.margin_top = Inches(0.25)
    
    p = tf_d.paragraphs[0]
    p.text = "🌐 Live Deployment & Access"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    
    dep_items = [
        ("Local Web Server", "Running locally at: http://127.0.0.1:5000 with real-time zero-cache headers."),
        ("Public Sharing (ngrok)", "Public HTTPS tunnel live for peer evaluation and recruiter demonstration."),
        ("Instant Startup", "Simple execution with 'python app.py' and 'ngrok http 5000'.")
    ]
    for title, desc in dep_items:
        p1 = tf_d.add_paragraph()
        p1.text = f"✔ {title}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_BLUE
        p2 = tf_d.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_LIGHT

    # Right box: Summary
    box_sum = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    box_sum.fill.solid()
    box_sum.fill.fore_color.rgb = CARD_BG
    box_sum.line.color.rgb = RGBColor(99, 165, 255)
    tf_s = box_sum.text_frame
    tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = Inches(0.25)
    
    p = tf_s.paragraphs[0]
    p.text = "🌟 Summary of Contributions"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    sum_items = [
        "Interactive glassmorphic web dashboard with responsive navigation.",
        "8-Model recommendation engine achieving 85% Precision@10.",
        "Modular syllabus breakdowns and direct companion book links.",
        "Diagnostic skill gap radar analysis comparing student vs employer thresholds.",
        "Embedded AI study advisor chatbot providing personalized learning guidance."
    ]
    for item in sum_items:
        p = tf_s.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    # Save presentation
    output_path = "CourseAI_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == '__main__':
    create_presentation()
